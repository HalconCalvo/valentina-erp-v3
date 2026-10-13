"""Migration Excel as the truth of the money of its OVs (advisor, 13/10): Folio_OV, download already filled,
sync with preview (add / change / cancel), never banks, OVs born in Valentina only report the total difference,
and the app does not capture money of an OV managed by the Excel."""
import io
from datetime import datetime

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlmodel import select

from app.core.config import settings
from app.models.audit import AuditFieldChange
from app.models.sales import (CustomerPayment, CustomerPaymentInstallment, CXCStatus, PaymentStatus, SalesOrder,
                              SalesOrderStatus)
from app.models.treasury import BankAccount, BankTransaction, TransactionType
from app.models.users import User
from app.schemas.sales_schema import PaymentPayload
from app.services import legacy_import_service as svc
from app.services import legacy_template_service as template
from app.services import sales_service
from tests.conftest import TEST_DIRECTOR_EMAIL
from tests.test_legacy_import_service import _inv, _messages, _ov, _workbook


@pytest.fixture
def director(session_fixture) -> User:
    return session_fixture.exec(select(User).where(User.email == TEST_DIRECTOR_EMAIL)).first()


def _two_invoices(session, director) -> SalesOrder:
    """Proyecto A: advance 101 paid in one payment and progress 102 with two payments."""
    data = _workbook([_ov()], [
        _inv(),
        _inv(Tipo="AVANCE", Folio_Factura=102, Fecha_Factura=datetime(2026, 2, 1), Monto_Factura=30000,
             Abono1_Fecha=datetime(2026, 2, 5), Abono1_Monto=10000, Abono2_Fecha=datetime(2026, 2, 20),
             Abono2_Monto=5000),
    ])
    assert svc.import_legacy_workbook(session, data, director)["errors"] == []
    return session.exec(select(SalesOrder)).first()


def _sheet_rows(wb, sheet: str) -> list[dict]:
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, [c for c in row])) for row in ws.iter_rows(min_row=2)]


def _save(wb) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _template(session):
    return load_workbook(io.BytesIO(template.build_template(session)))


def _active(session, model, **filters):
    rows = session.exec(select(model)).all()
    return [r for r in rows if all(getattr(r, k) == v for k, v in filters.items())]


def test_downloaded_template_is_filled_and_uploading_it_changes_nothing(session_fixture, director):
    order = _two_invoices(session_fixture, director)
    wb = _template(session_fixture)
    assert [c.value for c in wb["OVs"][1]][:2] == ["Folio_OV", "Proyecto"]
    assert [c.value for c in wb["Facturas"][1]][0] == "Folio_OV"
    invoices = _sheet_rows(wb, "Facturas")
    assert [r["Folio_OV"].value for r in invoices] == [f"OV-{order.id:04d}"] * 2
    assert invoices[1]["Abono2_Monto"].value == 5000
    preview = svc.validate_legacy_workbook(session_fixture, _save(wb))
    assert preview["errors"] == []
    assert preview["orders"][0]["action"] == "SIN CAMBIOS"
    assert preview["can_import"] is False


def test_sync_adds_changes_and_cancels_with_preview_and_change_log(session_fixture, director):
    order = _two_invoices(session_fixture, director)
    wb = _template(session_fixture)
    rows = _sheet_rows(wb, "Facturas")
    rows[1]["Abono2_Monto"].value = 7000          # changed payment
    rows[0]["Folio_OV"].value = None              # advance 101 no longer comes → cancelled
    rows[0]["Proyecto"].value = None
    for cell in rows[0].values():
        cell.value = None
    wb["Facturas"].append([f"OV-{order.id:04d}", "Proyecto A", "AVANCE", 103, datetime(2026, 3, 1), 20000,
                           None, 0, None, 0, None, datetime(2026, 3, 2), 20000])
    data = _save(wb)
    preview = svc.validate_legacy_workbook(session_fixture, data)
    assert preview["errors"] == [], preview["errors"]
    synced = preview["orders"][0]
    assert (synced["action"], synced["adds"], synced["changes"], synced["cancels"]) == ("SINCRONIZAR", 1, 1, 1)
    text = " | ".join(synced["details"])
    assert "Agrega factura 103 AVANCE por $20,000.00 con 1 abono(s)" in text
    assert "Factura 102 AVANCE: abono del 20/02/2026 por $5,000.00 → abono del 20/02/2026 por $7,000.00" in text
    assert "Factura 101 ANTICIPO ($58,000.00, 1 abono(s)) ya no viene en el Excel: se cancela" in text

    result = svc.import_legacy_workbook(session_fixture, data, director)
    assert (result["orders_synced"], result["changes_applied"]) == (1, 3)
    session_fixture.refresh(order)
    advance = session_fixture.exec(select(CustomerPayment).where(CustomerPayment.invoice_folio == "101")).one()
    assert advance.status == CXCStatus.CANCELLED
    assert all(i.is_cancelled for i in _active(session_fixture, CustomerPaymentInstallment,
                                                customer_payment_id=advance.id))
    # Balance = 116,000 − (10,000 + 7,000 + 20,000)
    assert order.outstanding_balance == 79000.0 and order.payment_status == PaymentStatus.PARTIAL
    reasons = {c.reason for c in session_fixture.exec(select(AuditFieldChange)).all() if c.reason}
    assert any(r.startswith("Importación ") for r in reasons)


def test_invoice_without_folio_matches_by_type_date_and_amount(session_fixture, director):
    order = _two_invoices(session_fixture, director)
    stored = session_fixture.exec(select(CustomerPayment).where(CustomerPayment.invoice_folio == "102")).one()
    stored.invoice_folio = None
    session_fixture.add(stored)
    session_fixture.commit()
    wb = _template(session_fixture)
    _sheet_rows(wb, "Facturas")[1]["Folio_Factura"].value = 492
    preview = svc.validate_legacy_workbook(session_fixture, _save(wb))
    assert preview["orders"][0]["details"] == ["Factura sin folio AVANCE $30,000.00: Folio — → 492"]
    svc.import_legacy_workbook(session_fixture, _save(wb), director)
    session_fixture.refresh(stored)
    assert stored.invoice_folio == "492" and stored.status != CXCStatus.CANCELLED
    assert order.id == stored.sales_order_id


def test_payment_linked_to_the_bank_is_never_touched(session_fixture, director):
    _two_invoices(session_fixture, director)
    account = BankAccount(name="Banco", account_number="1", current_balance=0.0, is_active=True)
    session_fixture.add(account)
    session_fixture.commit()
    movement = BankTransaction(account_id=account.id, transaction_type=TransactionType.IN, amount=5000.0,
                               description="Cobro")
    session_fixture.add(movement)
    session_fixture.commit()
    payment = session_fixture.exec(select(CustomerPaymentInstallment)
                                   .where(CustomerPaymentInstallment.amount == 5000.0)).one()
    payment.bank_transaction_id = movement.id
    session_fixture.add(payment)
    session_fixture.commit()
    wb = _template(session_fixture)
    row = _sheet_rows(wb, "Facturas")[1]
    row["Abono2_Fecha"].value, row["Abono2_Monto"].value = None, None
    preview = svc.validate_legacy_workbook(session_fixture, _save(wb))
    assert "está ligado a un movimiento de banco" in _messages(preview["errors"])
    assert preview["can_import"] is False


def test_order_born_in_valentina_only_reports_the_total_difference(session_fixture, director, seed_client_and_tax):
    client, tax = seed_client_and_tax
    order = SalesOrder(client_id=client.id, tax_rate_id=tax.id, project_name="Nacida en Valentina",
                       status=SalesOrderStatus.SOLD, valid_until=datetime(2026, 12, 31), subtotal=1000.0,
                       tax_amount=160.0, total_price=1160.0, outstanding_balance=1160.0)
    session_fixture.add(order)
    session_fixture.commit()
    folio = f"OV-{order.id:04d}"
    data = _workbook([_ov(Folio_OV=folio, Proyecto="Nacida en Valentina", Total_OV_con_IVA=2000)],
                     [_inv(Folio_OV=folio, Proyecto="Nacida en Valentina", Monto_Factura=500,
                           Abono1_Monto=500)])
    preview = svc.validate_legacy_workbook(session_fixture, data)
    assert preview["errors"] == []
    assert "no se cambia; los cambios de total van por orden de cambio" in _messages(preview["warnings"])
    svc.import_legacy_workbook(session_fixture, data, director)
    session_fixture.refresh(order)
    assert order.total_price == 1160.0 and order.outstanding_balance == 660.0 and order.is_excel_managed


def test_unknown_or_cancelled_folio_is_an_error(session_fixture, director):
    order = _two_invoices(session_fixture, director)
    order.status = SalesOrderStatus.CANCELLED_OV
    session_fixture.add(order)
    session_fixture.commit()
    data = _workbook([_ov(Folio_OV=f"OV-{order.id:04d}"), _ov(Folio_OV="OV-9999", Proyecto="X")], [])
    text = _messages(svc.validate_legacy_workbook(session_fixture, data)["errors"])
    assert f"OV-{order.id:04d} está cancelada." in text and "OV-9999 no existe en Valentina." in text


def test_order_missing_from_the_excel_goes_back_to_the_app(session_fixture, director):
    order = _two_invoices(session_fixture, director)
    data = _workbook([_ov(Proyecto="Proyecto Nuevo")], [])
    preview = svc.validate_legacy_workbook(session_fixture, data)
    assert preview["orders_released"] == 1
    svc.import_legacy_workbook(session_fixture, data, director)
    session_fixture.refresh(order)
    assert order.is_excel_managed is False
    assert all(c.status != CXCStatus.CANCELLED for c in _active(session_fixture, CustomerPayment,
                                                                 sales_order_id=order.id))


def test_app_does_not_capture_money_of_an_order_managed_by_the_excel(session_fixture, director):
    order = _two_invoices(session_fixture, director)
    with pytest.raises(HTTPException) as exc:
        sales_service.emit_advance_invoice(session_fixture, order.id,
                                           PaymentPayload(amount=100.0, invoice_folio="900"), director)
    assert exc.value.status_code == 409 and "se maneja por el Excel de migración" in exc.value.detail


def test_template_endpoint_requires_director_or_manager(client_fixture, sales_token, auth_header_director):
    url = f"{settings.API_V1_STR}/sales/orders/legacy-import/template"
    assert client_fixture.get(url, headers={"Authorization": f"Bearer {sales_token}"}).status_code == 403
    response = client_fixture.get(url, headers=auth_header_director)
    assert response.status_code == 200
    assert load_workbook(io.BytesIO(response.content)).sheetnames == ["OVs", "Facturas", "Instrucciones"]


def test_retention_follows_the_fg_note_only_when_the_note_changes(session_fixture, director):
    data = _workbook([_ov(Anticipo_Cobrado=None)], [_inv(Tipo="AVANCE", NC_FG_Folio=7, NC_FG_Monto=5800,
                                                         Abono1_Fecha=None, Abono1_Monto=None)])
    svc.import_legacy_workbook(session_fixture, data, director)
    cxc = session_fixture.exec(select(CustomerPayment)).one()
    cxc.retention_amount, cxc.retention_status = 0.0, None  # stored without retention (older import)
    session_fixture.add(cxc)
    session_fixture.commit()
    unchanged = svc.validate_legacy_workbook(session_fixture, _save(_template(session_fixture)))
    assert unchanged["orders"][0]["action"] == "SIN CAMBIOS"
    wb = _template(session_fixture)
    _sheet_rows(wb, "Facturas")[0]["NC_FG_Monto"].value = 2900
    details = svc.validate_legacy_workbook(session_fixture, _save(wb))["orders"][0]["details"]
    assert "Monto NC FG $5,800.00 → $2,900.00" in details[0] and "Retención $0.00 → $2,900.00" in details[0]
