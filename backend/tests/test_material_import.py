"""Sanitation tool 4: bulk update of the material catalog from Excel (docs/SANEAMIENTO.md §5.3)."""
import io

from openpyxl import Workbook, load_workbook
from sqlmodel import select

from app.core.config import settings
from app.core.security import create_access_token, get_password_hash
from app.models.audit import AuditFieldChange
from app.models.foundations import Provider
from app.models.material import Material, ProductionRoute
from app.models.users import User, UserRole
from app.services.material_import_service import COLUMNS

API = f"{settings.API_V1_STR}/material-import"


def _headers(session, role: UserRole) -> dict:
    user = User(email=f"{role.value.lower()}@imp.local", full_name=role.value, role=role, is_active=True,
                hashed_password=get_password_hash("Pass123!"))
    session.add(user)
    session.commit()
    session.refresh(user)
    return {"Authorization": f"Bearer {create_access_token(subject=user.email, user_id=user.id, user_role=role.value)}"}


def _seed(session) -> Provider:
    provider = Provider(business_name="Herrajes del Sureste", credit_days=0, is_active=True)
    session.add(provider)
    session.commit()
    session.add(Material(sku="IMP-1", name="Bisagra", category="HERRAJE", production_route=ProductionRoute.MATERIAL,
                         purchase_unit="", usage_unit="", current_cost=0.0))
    session.add(Material(sku="IMP-2", name="Tornillo", category="HERRAJE", production_route=ProductionRoute.MATERIAL,
                         purchase_unit="Caja", usage_unit="Pza", conversion_factor=100, current_cost=50.0,
                         provider_id=provider.id))
    session.commit()
    return provider


def _xlsx(rows: list) -> dict:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Materiales"
    sheet.append([header for _, header, _ in COLUMNS])
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return {"file": ("materiales.xlsx", output.getvalue(), "application/octet-stream")}


def test_validate_shows_changes_and_errors(client_fixture, session_fixture, auth_header_director):
    _seed(session_fixture)
    files = _xlsx([["IMP-1", "Bisagra", "herrajes del sureste", "Pza", "Pza", None, 35.5],
                   ["NO-EXISTE", None, None, "Pza", None, None, None],
                   ["IMP-1", None, None, None, None, None, None],
                   ["IMP-2", None, "Proveedor X", None, None, 0, None]])
    preview = client_fixture.post(f"{API}/validate", headers=auth_header_director, files=files).json()
    assert [(r["sku"], [c["field"] for c in r["changes"]]) for r in preview["rows"]] == [
        ("IMP-1", ["provider_id", "purchase_unit", "usage_unit", "current_cost"])]
    assert preview["rows"][0]["changes"][0]["new"] == "Herrajes del Sureste"
    assert [(e["row"], e["message"].split(" ")[0] + " " + e["message"].split(" ")[1]) for e in preview["errors"]] == [
        (3, "SKU vacío"), (4, "SKU repetido"), (5, "Proveedor 'Proveedor")]


def test_price_needs_price_role_and_apply_is_all_or_nothing(client_fixture, session_fixture, auth_header_director):
    _seed(session_fixture)
    warehouse = _headers(session_fixture, UserRole.WAREHOUSE)
    priced = _xlsx([["IMP-1", None, None, "Pza", "Pza", None, 35.5]])
    errors = client_fixture.post(f"{API}/validate", headers=warehouse, files=priced).json()["errors"]
    assert errors[0]["message"].startswith("Solo Dirección")
    units_only = _xlsx([["IMP-1", None, None, "Pza", "Pza", None, None]])
    assert client_fixture.post(f"{API}/validate", headers=warehouse, files=units_only).json()["rows"]
    assert client_fixture.post(f"{API}/validate", headers=_headers(session_fixture, UserRole.SALES),
                               files=units_only).status_code == 403
    rejected = client_fixture.post(f"{API}/apply", headers=warehouse, files=priced, data={"reason": "Catálogo"})
    assert rejected.status_code == 422
    session_fixture.expire_all()
    assert session_fixture.exec(select(Material).where(Material.sku == "IMP-1")).first().purchase_unit == ""


def test_apply_updates_with_reason_in_the_change_log(client_fixture, session_fixture, auth_header_director):
    provider = _seed(session_fixture)
    files = _xlsx([["IMP-1", None, provider.id, "Pza", "Pza", None, 35.5]])
    assert client_fixture.post(f"{API}/apply", headers=auth_header_director, files=files,
                               data={"reason": " "}).status_code == 422
    result = client_fixture.post(f"{API}/apply", headers=auth_header_director, files=files,
                                 data={"reason": "Catálogo completado por compras"})
    assert result.status_code == 200, result.text
    assert result.json() == {"updated": 1}
    session_fixture.expire_all()
    material = session_fixture.exec(select(Material).where(Material.sku == "IMP-1")).first()
    assert (material.provider_id, material.purchase_unit, material.current_cost) == (provider.id, "Pza", 35.5)
    logged = session_fixture.exec(select(AuditFieldChange).where(AuditFieldChange.table_name == "materials")).all()
    assert any((row.reason or "").startswith("Importación masiva del catálogo: Catálogo completado") for row in logged)


def test_template_lists_incomplete_materials(client_fixture, session_fixture, auth_header_director):
    _seed(session_fixture)
    response = client_fixture.get(f"{API}/template", headers=auth_header_director)
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content))["Materiales"]
    assert [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["IMP-1"]
    files = {"file": ("plantilla.xlsx", response.content, "application/octet-stream")}
    unchanged = client_fixture.post(f"{API}/validate", headers=auth_header_director, files=files).json()
    assert (unchanged["rows"], unchanged["errors"], unchanged["unchanged"]) == ([], [], 1)
