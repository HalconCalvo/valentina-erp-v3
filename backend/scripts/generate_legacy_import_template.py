"""Genera frontend/src/assets/plantilla-migracion-ov.xlsx (ejecutar una vez)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "frontend" / "src" / "assets" / "plantilla-migracion-ov.xlsx"

GRAY = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
ITALIC = Font(italic=True, color="64748B")

wb = Workbook()
ov = wb.active
ov.title = "OVs"
ov_headers = [
    "Proyecto",
    "Cliente",
    "Vendedor",
    "IVA%",
    "Total_OV_con_IVA",
    "Anticipo_Pct",
    "Anticipo_Cobrado",
    "Comision_Pct",
    "Notas",
]
ov.append(ov_headers)
ov.append(
    [
        "# EJEMPLO Residencial Las Palmas",
        "Cliente Demo SA de CV",
        "María Vendedora",
        16,
        1160000,
        60,
        696000,
        3,
        "Migración corte marzo 2026",
    ]
)
ov.append(
    [
        "# EJEMPLO Torre Centro",
        "Constructora Norte",
        "Juan Pérez",
        16,
        580000,
        50,
        290000,
        2.5,
        "Solo saldos abiertos",
    ]
)
for row in ov.iter_rows(min_row=2, max_row=3):
    for cell in row:
        cell.fill = GRAY
        cell.font = ITALIC

inv = wb.create_sheet("Facturas")
inv_headers = [
    "Proyecto",
    "Tipo",
    "Folio_Factura",
    "Fecha_Factura",
    "Monto_Factura",
    "NC_Anticipo_Folio",
    "NC_Anticipo_Monto",
    "NC_FG_Folio",
    "NC_FG_Monto",
    "Abono1_Fecha",
    "Abono1_Monto",
    "Abono2_Fecha",
    "Abono2_Monto",
    "Abono3_Fecha",
    "Abono3_Monto",
]
inv.append(inv_headers)
inv.append(
    [
        "# EJEMPLO Residencial Las Palmas",
        "ANTICIPO",
        "FA-1001",
        "2025-11-15",
        696000,
        "",
        0,
        "",
        0,
        "2025-11-20",
        696000,
        "",
        0,
        "",
        0,
    ]
)
inv.append(
    [
        "# EJEMPLO Residencial Las Palmas",
        "AVANCE",
        "FA-1002",
        "2026-01-10",
        300000,
        "NC-A-01",
        5000,
        "",
        0,
        "2026-01-15",
        150000,
        "2026-02-01",
        100000,
        "",
        0,
    ]
)
inv.append(
    [
        "# EJEMPLO Torre Centro",
        "CONTRATO",
        "FA-2001",
        "2025-12-01",
        580000,
        "",
        0,
        "NC-FG-01",
        29000,
        "2025-12-05",
        290000,
        "",
        0,
        "",
        0,
    ]
)
for row in inv.iter_rows(min_row=2, max_row=4):
    for cell in row:
        cell.fill = GRAY
        cell.font = ITALIC

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
