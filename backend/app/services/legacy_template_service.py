"""Migration Excel downloaded from the app (advisor, 13/10): already filled with the OVs that came in the Excel, their
invoices, credit notes and payments, in exactly the format the importer reads (Folio_OV first). Uploading it back
without changes changes nothing. Also the downloadable report of differences (D26).
"""
import io
from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlmodel import Session

from app.core.business_time import utc_naive_to_local
from app.models.foundations import Client, TaxRate
from app.models.sales import CustomerPayment, SalesOrder
from app.models.users import User
from app.repositories import legacy_import_repository as legacy_repo
from app.services import legacy_import_service as importer
from app.services import legacy_sync_service as legacy_sync

DATE_FORMAT = "DD/MM/YYYY"
MONEY_FORMAT = "#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
NOTES = [
    "Cómo se usa (Gerencia / Dirección):",
    "1. Este archivo es la verdad del DINERO de las OVs que contiene: facturas, notas de crédito, anticipos y abonos.",
    "2. Folio_OV liga cada factura con su OV. Déjalo vacío solo para una OV nueva (sus facturas se ligan por Proyecto).",
    "3. La subida solo AGREGA lo que falta en Valentina (OVs nuevas, facturas y abonos). Nunca cambia ni cancela.",
    "4. Las diferencias (montos o fechas distintos, registros que están de un lado y no del otro) se reportan y se",
    "   corrigen a mano en la app con motivo. Reporte descargable en la misma pantalla. No toca bancos ni comisiones.",
    "5. Todo se sigue capturando en la app. Una OV que no viene en el archivo no se toca: se marca para revisar.",
    "6. Costo_OV (opcional, sin IVA): el costo de la OV legacy. Con él, la OV entra al Margen neto % sobre venta del",
    "   tablero de Dirección; vacío, queda fuera. Cambiar un costo ya capturado se reporta y se aplica con motivo.",
]


def _tax_label(rate: Optional[TaxRate]) -> Any:
    if rate is None:
        return None
    if importer._is_exempt_rate(rate):  # noqa: SLF001
        return importer.EXEMPT_LABEL
    return round(float(rate.rate) * 100, 4)


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    ws.freeze_panes = "B2"


def _format_columns(ws, headers: list[str]) -> None:
    for index, name in enumerate(headers, start=1):
        letter = ws.cell(row=1, column=index).column_letter
        ws.column_dimensions[letter].width = max(12, len(name) + 2)
        number_format = DATE_FORMAT if "Fecha" in name else MONEY_FORMAT if "Monto" in name or "Total" in name \
            or name in ("Anticipo_Cobrado", importer.COST_COLUMN) else None
        if number_format:
            for cell in ws[letter][1:]:
                cell.number_format = number_format


def _order_row(session: Session, order: SalesOrder, names: dict, rates: dict,
               invoices: list[CustomerPayment]) -> list:
    advance = sum(float(c.amount or 0.0) for c in invoices if c.payment_type.value == "ADVANCE")
    cost = legacy_sync.order_cost(legacy_repo.get_active_items(session, order.id))
    return [importer.order_folio(order.id), order.project_name, names["clients"].get(order.client_id),
            names["users"].get(order.user_id), _tax_label(rates.get(order.tax_rate_id)),
            round(float(order.total_price or 0.0), 2), round(advance, 2),
            float(order.applied_commission_percent or 0.0), order.notes, cost if cost > 0 else None]


def _invoice_row(session: Session, order: SalesOrder, cxc: CustomerPayment, slots: int) -> list:
    has_fg = bool(cxc.nc_retention_folio) and float(cxc.nc_retention_amount or 0.0) > 0
    row = [importer.order_folio(order.id), order.project_name, legacy_sync.TYPE_LABELS.get(cxc.payment_type),
           cxc.invoice_folio or None, cxc.invoice_date, float(cxc.amount or 0.0), cxc.nc_advance_folio or None,
           float(cxc.nc_advance_amount or 0.0), cxc.nc_retention_folio or None, float(cxc.nc_retention_amount or 0.0),
           cxc.retention_due_date if has_fg else None]
    installments = legacy_repo.get_active_installments(session, cxc.id)
    for index in range(slots):
        inst = installments[index] if index < len(installments) else None
        row += [inst.payment_date, float(inst.amount or 0.0)] if inst else [None, None]
    return row


def _invoice_headers(slots: int) -> list[str]:
    base = [h for h in importer.INV_HEADERS if not h.startswith("Abono")]
    fg_index = base.index("NC_FG_Monto") + 1
    headers = base[:fg_index] + ["FG_Fecha_Vencimiento"] + base[fg_index:]
    return headers + [f"Abono{n}_{part}" for n in range(1, slots + 1) for part in ("Fecha", "Monto")]


def _lookups(session: Session, orders: list[SalesOrder]) -> tuple[dict, dict]:
    names = {"clients": legacy_repo.get_full_names(session, Client, [o.client_id for o in orders]),
             "users": legacy_repo.get_full_names(session, User, [o.user_id for o in orders if o.user_id])}
    return names, legacy_repo.get_tax_rates_by_id(session, list({o.tax_rate_id for o in orders}))


def build_template(session: Session) -> bytes:
    orders = legacy_repo.get_excel_managed_orders(session)
    invoices = {o.id: legacy_repo.get_active_invoices(session, o.id) for o in orders}
    most = max([len(legacy_repo.get_active_installments(session, c.id)) for cs in invoices.values() for c in cs],
               default=0)
    slots = max(3, min(most, importer.MAX_INSTALLMENT_SLOTS))
    names, rates = _lookups(session, orders)
    wb = Workbook()
    ws = wb.active
    ws.title = importer.OV_SHEET
    ov_headers = importer.OV_HEADERS + importer.OV_OPTIONAL_HEADERS
    _write_header(ws, ov_headers)
    for order in orders:
        ws.append(_order_row(session, order, names, rates, invoices[order.id]))
    _format_columns(ws, ov_headers)
    ws_inv = wb.create_sheet(importer.INV_SHEET)
    headers = _invoice_headers(slots)
    _write_header(ws_inv, headers)
    for order in orders:
        for cxc in invoices[order.id]:
            ws_inv.append(_invoice_row(session, order, cxc, slots))
    _format_columns(ws_inv, headers)
    notes = wb.create_sheet("Instrucciones")
    for line in NOTES + ["", f"Descargado el {utc_naive_to_local(datetime.utcnow()):%d/%m/%Y %H:%M} (hora de Mérida)."]:
        notes.append([line])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


REPORT_HEADERS = ["Folio OV", "Proyecto", "Diferencia", "Factura / abono", "Concepto", "En Valentina", "En el Excel",
                  "Corrección (motivo)"]


def build_discrepancy_report(rows: list[dict]) -> bytes:
    """Differences between Valentina and the Excel, to correct by hand in the app with a reason (D26)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Diferencias"
    _write_header(ws, REPORT_HEADERS)
    for row in rows:
        ws.append([row["folio"], row["project"], row["kind"], row["record"], row["concept"], row["valentina"],
                   row["excel"], None])
    for letter, width in zip("ABCDEFGH", (11, 30, 30, 34, 20, 34, 34, 40)):
        ws.column_dimensions[letter].width = width
    ws.append([])
    ws.append([f"Generado el {utc_naive_to_local(datetime.utcnow()):%d/%m/%Y %H:%M} (hora de Mérida). "
               "La subida no cambió ni canceló nada de esto."])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
