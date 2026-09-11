"""Importación de OVs legacy desde Excel (corte en ceros)."""
from __future__ import annotations

import io
from datetime import datetime, timedelta
from typing import Any

import structlog
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlmodel import Session, select

from app.models.foundations import Client, TaxRate
from app.models.sales import (
    CXCStatus,
    CustomerPayment,
    CustomerPaymentInstallment,
    PaymentMethod,
    PaymentStatus,
    PaymentType,
    SalesOrder,
    SalesOrderItem,
    SalesOrderStatus,
)
from app.models.users import User
from app.services import sales_service

logger = structlog.get_logger(__name__)

OV_HEADERS = [
    "Proyecto",
    "Cliente",
    "Vendedor",
    "IVA%",
    "Total_OV_con_IVA",
    "Anticipo_Pct",
    "Anticipo_Cobrado",
    "Comision_Pct",
    "Notas",
]
INV_HEADERS = [
    "Proyecto",
    "Tipo",
    "Folio_Factura",
    "Fecha_Factura",
    "Monto_Factura",
    "NC_Anticipo_Folio",
    "NC_Anticipo_Monto",
    "NC_FG_Folio",
    "NC_FG_Monto",
    "Abono1_Fecha",
    "Abono1_Monto",
    "Abono2_Fecha",
    "Abono2_Monto",
    "Abono3_Fecha",
    "Abono3_Monto",
]
TIPO_MAP = {
    "ANTICIPO": PaymentType.ADVANCE,
    "AVANCE": PaymentType.PROGRESS,
    "CONTRATO": PaymentType.FULL,
}


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _parse_date(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    text = _cell_str(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _read_sheet_rows(wb, sheet_name: str, headers: list[str]) -> list[dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"Falta la hoja '{sheet_name}'.")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header_row = [_cell_str(c) for c in rows[0]]
    idx = {name: header_row.index(name) for name in headers if name in header_row}
    missing = [h for h in headers if h not in idx]
    if missing:
        raise HTTPException(status_code=400, detail=f"Columnas faltantes en {sheet_name}: {', '.join(missing)}")
    out: list[dict[str, Any]] = []
    for raw in rows[1:]:
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue
        record = {h: raw[idx[h]] if idx[h] < len(raw) else None for h in headers}
        proyecto = _cell_str(record.get("Proyecto"))
        if not proyecto or proyecto.startswith("#"):
            continue
        out.append(record)
    return out


def _find_client(session: Session, name: str) -> Client | None:
    target = name.strip()
    return session.exec(select(Client).where(Client.full_name == target)).first()


def _find_seller(session: Session, name: str) -> User | None:
    target = name.strip()
    if not target:
        return None
    return session.exec(select(User).where(User.full_name == target)).first()


def _find_tax_rate(session: Session, iva_percent: float) -> TaxRate:
    rate = iva_percent / 100.0 if iva_percent > 1 else iva_percent
    tax = session.exec(
        select(TaxRate).where(TaxRate.is_active == True)  # noqa: E712
    ).all()
    for t in tax:
        if abs(float(t.rate) - rate) < 0.0001:
            return t
    default = session.exec(select(TaxRate).where(TaxRate.is_active == True)).first()  # noqa: E712
    if not default:
        raise HTTPException(status_code=400, detail="No hay tasas de IVA configuradas.")
    return default


def _legacy_order_exists(session: Session, project_name: str) -> bool:
    row = session.exec(
        select(SalesOrder).where(
            SalesOrder.project_name == project_name,
            SalesOrder.is_legacy == True,  # noqa: E712
        )
    ).first()
    return row is not None


def _map_payment_type(tipo: str) -> PaymentType:
    key = _cell_str(tipo).upper()
    if key not in TIPO_MAP:
        raise ValueError(f"Tipo de factura inválido: {tipo}")
    return TIPO_MAP[key]


def _create_legacy_order(
    session: Session,
    ov: dict[str, Any],
    client: Client,
    seller: User | None,
    tax_rate: TaxRate,
    director_id: int,
) -> SalesOrder:
    total_price = round(_cell_float(ov.get("Total_OV_con_IVA")), 2)
    iva_pct = _cell_float(ov.get("IVA%"), 16.0)
    rate = iva_pct / 100.0 if iva_pct > 1 else iva_pct
    subtotal = round(total_price / (1.0 + rate), 2) if rate >= 0 else total_price
    tax_amount = round(total_price - subtotal, 2)
    commission_pct = _cell_float(ov.get("Comision_Pct"), 0.0)
    commission_amount = round(subtotal * (commission_pct / 100.0), 2)

    order = SalesOrder(
        client_id=client.id,
        tax_rate_id=tax_rate.id,
        user_id=seller.id if seller else director_id,
        project_name=_cell_str(ov.get("Proyecto")),
        status=SalesOrderStatus.SOLD,
        is_legacy=True,
        is_approved_by_director=True,
        director_approved_at=datetime.utcnow(),
        valid_until=datetime.utcnow() + timedelta(days=365),
        advance_percent=_cell_float(ov.get("Anticipo_Pct"), 60.0),
        advance_invoice_amount=_cell_float(ov.get("Anticipo_Cobrado")) or None,
        has_advance_invoice=_cell_float(ov.get("Anticipo_Cobrado")) > 0,
        applied_commission_percent=commission_pct,
        commission_amount=commission_amount,
        subtotal=subtotal,
        tax_amount=tax_amount,
        total_price=total_price,
        outstanding_balance=total_price,
        payment_status=PaymentStatus.PENDING,
        notes=_cell_str(ov.get("Notas")) or "Importación legacy",
    )
    session.add(order)
    session.flush()

    item = SalesOrderItem(
        sales_order_id=order.id,
        product_name="Contrato legacy (migración)",
        quantity=1.0,
        unit_price=subtotal,
        subtotal_price=subtotal,
        cost_snapshot={"legacy": True},
        frozen_unit_cost=0.0,
    )
    session.add(item)
    session.flush()
    return order


def _add_invoice_and_installments(
    session: Session,
    order: SalesOrder,
    inv: dict[str, Any],
    director_id: int,
) -> tuple[int, int]:
    ptype = _map_payment_type(inv.get("Tipo"))
    amount = round(_cell_float(inv.get("Monto_Factura")), 2)
    if amount <= 0:
        raise ValueError("Monto_Factura debe ser mayor a cero.")

    cxc = CustomerPayment(
        sales_order_id=order.id,
        payment_type=ptype,
        invoice_folio=_cell_str(inv.get("Folio_Factura")) or None,
        invoice_date=_parse_date(inv.get("Fecha_Factura")) or datetime.utcnow(),
        amount=amount,
        status=CXCStatus.PENDING,
        payment_method=PaymentMethod.TRANSFER,
        created_by_user_id=director_id,
        nc_advance_folio=_cell_str(inv.get("NC_Anticipo_Folio")) or None,
        nc_advance_amount=round(_cell_float(inv.get("NC_Anticipo_Monto")), 2),
        nc_retention_folio=_cell_str(inv.get("NC_FG_Folio")) or None,
        nc_retention_amount=round(_cell_float(inv.get("NC_FG_Monto")), 2),
    )
    session.add(cxc)
    session.flush()

    installments = 0
    for i in range(1, 4):
        monto = round(_cell_float(inv.get(f"Abono{i}_Monto")), 2)
        if monto <= 0:
            continue
        fecha = _parse_date(inv.get(f"Abono{i}_Fecha")) or datetime.utcnow()
        inst = CustomerPaymentInstallment(
            customer_payment_id=cxc.id,
            amount=monto,
            payment_date=fecha,
            payment_method=PaymentMethod.TRANSFER,
            created_by_user_id=director_id,
            is_advance=ptype == PaymentType.ADVANCE,
        )
        session.add(inst)
        order.outstanding_balance = float(order.outstanding_balance or 0.0) - monto
        installments += 1

    nc_credit = float(cxc.nc_advance_amount or 0.0) + float(cxc.nc_retention_amount or 0.0)
    if nc_credit > 0:
        order.outstanding_balance = float(order.outstanding_balance or 0.0) - nc_credit

    sales_service._recalc_cxc_status(session, cxc, order)  # noqa: SLF001
    session.add(cxc)
    session.add(order)
    return 1, installments


def _finalize_order_payment_status(order: SalesOrder) -> None:
    if float(order.outstanding_balance or 0.0) <= 0.1:
        order.outstanding_balance = 0.0
        order.payment_status = PaymentStatus.PAID
    else:
        order.payment_status = PaymentStatus.PARTIAL


def import_legacy_workbook(session: Session, file_bytes: bytes, director: User) -> dict:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ov_rows = _read_sheet_rows(wb, "OVs", OV_HEADERS)
    inv_rows = _read_sheet_rows(wb, "Facturas", INV_HEADERS)
    invoices_by_project: dict[str, list[dict[str, Any]]] = {}
    for inv in inv_rows:
        key = _cell_str(inv.get("Proyecto"))
        invoices_by_project.setdefault(key, []).append(inv)

    result = {
        "orders_created": 0,
        "invoices_created": 0,
        "installments_created": 0,
        "orders_created_details": [],
        "warnings": [],
        "errors": [],
    }

    for ov in ov_rows:
        project = _cell_str(ov.get("Proyecto"))
        try:
            if _legacy_order_exists(session, project):
                result["warnings"].append(f"OV legacy ya existe, omitida: {project}")
                continue

            client_name = _cell_str(ov.get("Cliente"))
            client = _find_client(session, client_name)
            if not client:
                result["errors"].append(f"{project}: cliente no encontrado «{client_name}»")
                continue

            seller = _find_seller(session, _cell_str(ov.get("Vendedor")))
            tax_rate = _find_tax_rate(session, _cell_float(ov.get("IVA%"), 16.0))

            order = _create_legacy_order(
                session, ov, client, seller, tax_rate, director.id
            )
            inv_count = 0
            inst_count = 0
            for inv in invoices_by_project.get(project, []):
                added_inv, added_inst = _add_invoice_and_installments(
                    session, order, inv, director.id
                )
                inv_count += added_inv
                inst_count += added_inst

            _finalize_order_payment_status(order)
            session.add(order)
            session.commit()
            session.refresh(order)

            folio = sales_service._order_reference(order.id)  # noqa: SLF001
            result["orders_created"] += 1
            result["invoices_created"] += inv_count
            result["installments_created"] += inst_count
            result["orders_created_details"].append(
                {
                    "project_name": project,
                    "order_id": order.id,
                    "folio": folio,
                    "outstanding_balance": float(order.outstanding_balance or 0.0),
                }
            )
            logger.info("legacy_ov_imported", project=project, order_id=order.id)
        except Exception as exc:
            session.rollback()
            msg = f"{project}: {exc}"
            result["errors"].append(msg)
            logger.error("legacy_ov_import_failed", project=project, error=str(exc))

    return result
