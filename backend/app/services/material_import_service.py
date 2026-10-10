"""Sanitation tool 4 (docs/SANEAMIENTO.md §5.3): bulk update of the material catalog from an Excel file.

One sheet, one row per material identified by SKU. Blank cell = no change. Columns: provider (name or id), purchase
unit, usage unit, conversion factor and cost per purchase unit. Validate → preview (before → after) → apply all or
nothing with a reason; catalog roles, price and factor only DIRECTOR/MANAGER/ADMIN (same rules as the material form).
"""
import io
import unicodedata
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlmodel import Session

from app.core.audit_context import audit_reason
from app.models.material import Material
from app.repositories import material_import_repository as repo
from app.schemas.material_import_schema import (
    MaterialChangeRowRead, MaterialFieldChangeRead, MaterialImportIssueRead, MaterialImportPreviewRead,
    MaterialImportResultRead,
)
from app.services import material_service

SHEET = "Materiales"
COLUMNS = [  # (key, header, material field)
    ("sku", "SKU", None), ("name", "Nombre (referencia)", None), ("provider", "Proveedor", "provider_id"),
    ("purchase_unit", "Unidad de compra", "purchase_unit"), ("usage_unit", "Unidad de uso", "usage_unit"),
    ("conversion_factor", "Factor de conversión", "conversion_factor"),
    ("current_cost", "Costo por unidad de compra", "current_cost"),
]
LABELS = {field: header for _, header, field in COLUMNS if field}


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return " ".join(text.strip().lower().split())


HEADER_KEYS = {_norm(header): key for key, header, _ in COLUMNS}


def _blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _number(value: Any, label: str, minimum: float, strict: bool) -> float:
    try:
        number = float(str(value).replace(",", "").replace("$", "").strip()) if isinstance(value, str) else float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label}: '{value}' no es un número.") from None
    if number < minimum or (strict and number == minimum):
        raise ValueError(f"{label} debe ser mayor {'a' if strict else 'o igual a'} {minimum:g}.")
    return round(number, 6)


def _provider(value: Any, providers: list) -> int:
    text = str(value).strip()
    if text.isdigit() and any(p.id == int(text) for p in providers):
        return int(text)
    matches = [p for p in providers if _norm(p.business_name) == _norm(text)]
    if len(matches) != 1:
        raise ValueError(f"Proveedor '{text}' {'repetido' if matches else 'no existe'} en el catálogo.")
    return matches[0].id


def _read_rows(content: bytes) -> List[Tuple[int, Dict[str, Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — any parse failure means an unreadable file
        raise HTTPException(status_code=400, detail="El archivo no es un .xlsx válido.") from exc
    sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    keys = [HEADER_KEYS.get(_norm(cell)) for cell in next(rows, [])]
    if "sku" not in keys:
        raise HTTPException(status_code=400, detail="Falta la columna SKU (usa la plantilla).")
    result = []
    for number, values in enumerate(rows, start=2):
        data = {key: value for key, value in zip(keys, values) if key}
        if any(not _blank(v) for v in data.values()):
            result.append((number, data))
    return result


def _parsed_values(data: Dict[str, Any], providers: list) -> Dict[str, Any]:
    values: Dict[str, Any] = {}
    if not _blank(data.get("provider")):
        values["provider_id"] = _provider(data["provider"], providers)
    for field in ("purchase_unit", "usage_unit"):
        if not _blank(data.get(field)):
            values[field] = str(data[field]).strip()
    if not _blank(data.get("conversion_factor")):
        values["conversion_factor"] = _number(data["conversion_factor"], LABELS["conversion_factor"], 0, True)
    if not _blank(data.get("current_cost")):
        values["current_cost"] = _number(data["current_cost"], LABELS["current_cost"], 0, False)
    return values


def _display(field: str, value: Any, provider_names: Dict[int, str]) -> Optional[str]:
    if value is None or value == "":
        return None
    return provider_names.get(value, str(value)) if field == "provider_id" else str(value)


def _changes(material: Material, values: Dict[str, Any], provider_names: Dict[int, str]) -> List[MaterialFieldChangeRead]:
    changes = []
    for field, new in values.items():
        old = getattr(material, field)
        same = float(old or 0) == float(new) if isinstance(new, float) else (old or None) == new
        if not same:
            changes.append(MaterialFieldChangeRead(field=field, label=LABELS[field],
                                                   old=_display(field, old, provider_names),
                                                   new=_display(field, new, provider_names)))
    return changes


def _plan(session: Session, content: bytes, user) -> Tuple[MaterialImportPreviewRead, Dict[str, Dict[str, Any]]]:
    material_service.assert_catalog_role(user)
    materials, providers = repo.materials_by_sku(session), repo.providers(session)
    provider_names = {p.id: p.business_name for p in providers}
    preview, updates, seen = MaterialImportPreviewRead(), {}, set()
    for number, data in _read_rows(content):
        sku = str(data.get("sku") or "").strip()
        material = materials.get(sku.upper())
        try:
            if not material or sku.upper() in seen:
                raise ValueError("SKU repetido en el archivo." if material else "SKU vacío o inexistente en el catálogo.")
            seen.add(sku.upper())
            values = _parsed_values(data, providers)
            changes = _changes(material, values, provider_names)
            if any(c.field in material_service.PRICE_FIELDS for c in changes) and not material_service.can_change_price(user):
                raise ValueError("Solo Dirección, Gerencia o Administración cambian costo o factor de conversión.")
        except ValueError as exc:
            preview.errors.append(MaterialImportIssueRead(row=number, sku=sku or None, message=str(exc)))
            continue
        if not changes:
            preview.unchanged += 1
            continue
        preview.rows.append(MaterialChangeRowRead(row=number, sku=material.sku, name=material.name, changes=changes))
        updates[material.sku.upper()] = {c.field: values[c.field] for c in changes}
    return preview, updates


def validate(session: Session, content: bytes, user) -> MaterialImportPreviewRead:
    return _plan(session, content, user)[0]


def apply(session: Session, content: bytes, reason: str, user) -> MaterialImportResultRead:
    text = (reason or "").strip()
    if not text:
        raise HTTPException(status_code=422, detail="El motivo es obligatorio.")
    preview, updates = _plan(session, content, user)
    if preview.errors:
        raise HTTPException(status_code=422, detail="El archivo tiene errores: valídalo y corrígelo antes de aplicar.")
    materials = repo.materials_by_sku(session)
    with audit_reason(f"Importación masiva del catálogo: {text}"):
        for sku, values in updates.items():
            for field, value in values.items():
                setattr(materials[sku], field, value)
            session.add(materials[sku])
        session.commit()
    return MaterialImportResultRead(updated=len(updates))


def template(session: Session) -> bytes:
    """The sheet prefilled with the active materials that miss provider, units or cost; plus the provider list."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append([header for _, header, _ in COLUMNS])
    names = {p.id: p.business_name for p in repo.providers(session)}
    for m in repo.incomplete_active_materials(session):
        sheet.append([m.sku, m.name, names.get(m.provider_id), m.purchase_unit or None, m.usage_unit or None,
                      m.conversion_factor, m.current_cost])
    reference = workbook.create_sheet("Proveedores")
    reference.append(["Id", "Proveedor"])
    for provider_id, name in sorted(names.items(), key=lambda item: item[1] or ""):
        reference.append([provider_id, name])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
